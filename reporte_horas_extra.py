"""
Generador de Reporte de Horas Extra
Lee el reporte de Varias Obras y los reportes de obra (.xlsm) y genera un Excel
con una hoja Resumen y una hoja por reporte.

Columnas de cada hoja: Trabajador, Ajustero, Supervisor/Tecnico, Oficio, Sal. Dia,
Coste/Hora, horas extra por dia (una col por dia), Dias quincena trabajados,
Coste dias / ajuste, Coste H. Extra, Total H. Extra, % H. Extra,
Coste total quincena, Dias atrasados, Horas extras atrasadas, Atrasos/otros,
coste total (H. Extra + atrasos) y comentarios.
"""

import sys, re, json, os
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "pandas", "openpyxl"])
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Configuración de umbrales (editable desde la interfaz, persistida en disco)
# ---------------------------------------------------------------------------

CONFIG_HORAS_DIA_ALERTA  = 5      # celda de día con más de N horas -> rojo/negrita
CONFIG_PORCENTAJE_MINIMO = 0.20   # % mínimo H. Extra sobre coste total -> alerta fila

_CONFIG_PATH = os.path.join(os.path.expanduser("~"), ".reporte_horas_extra_config.json")

def cargar_configuracion():
    """Lee la configuración del disco si existe."""
    global CONFIG_HORAS_DIA_ALERTA, CONFIG_PORCENTAJE_MINIMO
    try:
        with open(_CONFIG_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        CONFIG_HORAS_DIA_ALERTA  = int(data.get("horas_dia", CONFIG_HORAS_DIA_ALERTA))
        CONFIG_PORCENTAJE_MINIMO = float(data.get("porcentaje", CONFIG_PORCENTAJE_MINIMO))
    except Exception:
        pass

def guardar_configuracion():
    """Guarda la configuración actual en disco."""
    try:
        with open(_CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump({"horas_dia":  CONFIG_HORAS_DIA_ALERTA,
                       "porcentaje": CONFIG_PORCENTAJE_MINIMO}, f)
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Lectura de archivos
# ---------------------------------------------------------------------------

def _num(v):
    """Convierte a float; vacio / texto / NaN -> 0."""
    v = pd.to_numeric(v, errors='coerce')
    return float(v) if pd.notna(v) else 0.0

def _txt(v):
    return str(v).strip() if pd.notna(v) else ''

def _norm(col):
    """Normaliza cabeceras: quita saltos de linea, _x000D_ y espacios dobles."""
    return ' '.join(str(col).replace('_x000D_', ' ').split()).upper()

def _fecha_key(fecha):
    """Clave de orden cronologico para fechas 'dd-mm-yyyy'."""
    try:
        return datetime.strptime(fecha, '%d-%m-%Y')
    except ValueError:
        return datetime.max


def repartir_atrasos(importe, sal_dia, coste_hora, observaciones):
    """
    Reparte el importe de ATRASOS/OTROS del reporte en (dias, horas) de forma
    que dias * Sal. Dia + horas * Coste/Hora == importe.
    Los dias se deducen de las observaciones ("domingo", "5 dias"); el resto
    del importe se convierte en horas extra. Si el texto no habla de dias ni
    de horas y el importe es multiplo exacto del salario diario, se toma como dias.
    """
    if not importe:
        return None, None
    obs  = (observaciones or '').lower()
    dias = obs.count('domingo') + sum(int(n) for n in re.findall(r'(\d+)\s*d[ií]as?\b', obs))
    if dias and sal_dia and abs(dias * sal_dia) > abs(importe):
        dias = 0
    menciona_horas = re.search(r'\d\s*h', obs)
    if not dias and not menciona_horas and sal_dia and abs(importe) % sal_dia == 0:
        dias = importe / sal_dia
    if importe < 0:
        dias = -abs(dias)
    resto = importe - dias * sal_dia
    horas = round(resto / coste_hora, 4) if coste_hora else 0
    if not coste_hora and resto:
        # No se puede expresar en horas: se deja todo como dias equivalentes
        return (importe / sal_dia if sal_dia else None), None
    return (dias or None), (horas or None)


def _worker(nombre, ajustero, tecnico, supervisor, oficio, obra,
            sal_dia, coste_hora, dias_trab, coste_total, atrasos, obs, days):
    dias_atr, horas_atr = repartir_atrasos(atrasos, sal_dia, coste_hora, obs)
    total_e = sum(days.values())
    return {
        'TRABAJADOR':     nombre,
        'AJUSTERO':       ajustero,
        'TECNICO':        tecnico,
        'SUPERVISOR':     supervisor,
        'OFICIO':         oficio,
        'OBRA':           obra,
        'SAL. DIA':       sal_dia,
        'COSTE HORA':     coste_hora,
        'DIAS':           dias_trab,
        'COSTE TOTAL':    coste_total,
        'ATRASOS':        atrasos,
        'DIAS ATR':       dias_atr,
        'HORAS ATR':      horas_atr,
        'COMENTARIOS':    obs,
        'TOTAL H. EXTRA': total_e,
        'COSTE H. EXTRA': total_e * coste_hora,
        '_days':          days,
    }


def _incluir(w):
    """En las hojas solo aparecen trabajadores con horas extra o con atrasos."""
    return w['TOTAL H. EXTRA'] > 0 or w['ATRASOS'] != 0


def _es_codigo_obra(v):
    try:
        return 10000 <= int(v) <= 99999
    except Exception:
        return False


def _read_formato_diario(path):
    """
    Formato del reporte de Varias Obras: fila 1 con las fechas, fila 2 con
    cabeceras y por cada dia un bloque ADMIN, AJUSTE, OBRA 1, EXTRA 1, OBRA 2, EXTRA 2.
    Devuelve (workers_todos, coste_total_todos).
    """
    raw = pd.read_excel(path, header=None)
    date_row, header_row = raw.iloc[0], raw.iloc[1]

    cols = {}
    for i, h in enumerate(header_row):
        if pd.notna(h):
            cols.setdefault(_norm(h), i)

    day_blocks = []
    last_date  = None
    i = 0
    while i < len(header_row):
        m = re.search(r'(\d{2}-\d{2}-\d{4})', _txt(date_row.iloc[i]))
        if m:
            last_date = m.group(1)
        if last_date and _norm(header_row.iloc[i]) == 'ADMIN':
            day_blocks.append({
                'fecha':  last_date,
                'admin':  i,     'ajuste': i + 1,
                'obra1':  i + 2, 'extra1': i + 3,
                'obra2':  i + 4, 'extra2': i + 5,
            })
            i += 6
        else:
            i += 1

    def get(row, name):
        c = cols.get(name)
        return row.iloc[c] if c is not None else None

    workers = []
    for r in range(2, len(raw)):
        row    = raw.iloc[r]
        nombre = _txt(get(row, 'OPERARIOS'))
        if not nombre or nombre.upper() in ('TOTAL', 'TOTALES'):
            continue

        day_hours  = {}   # fecha -> horas extra del dia
        obra_hours = {}   # codigo_obra -> horas (regulares + extras)
        obra_def   = get(row, 'OBRA')
        obra_def   = str(int(obra_def)) if _es_codigo_obra(obra_def) else None

        for b in day_blocks:
            dias_reg = _num(row.iloc[b['admin']]) + _num(row.iloc[b['ajuste']])
            if dias_reg > 0 and obra_def:
                obra_hours[obra_def] = obra_hours.get(obra_def, 0) + dias_reg * 8
            for oc, ec in ((b['obra1'], b['extra1']), (b['obra2'], b['extra2'])):
                extra = _num(row.iloc[ec])
                if extra <= 0:
                    continue
                day_hours[b['fecha']] = day_hours.get(b['fecha'], 0) + extra
                if _es_codigo_obra(row.iloc[oc]):
                    cod = str(int(row.iloc[oc]))
                    obra_hours[cod] = obra_hours.get(cod, 0) + extra

        top = sorted(obra_hours.items(), key=lambda x: -x[1])[:2]

        workers.append(_worker(
            nombre      = nombre,
            ajustero    = _txt(get(row, 'AJUSTERO')).rstrip('*').strip(),
            tecnico     = _txt(get(row, 'TECNICO OBRA')),
            supervisor  = '',
            oficio      = _txt(get(row, 'OFICIO')),
            obra        = '/'.join(cod for cod, _ in top),
            sal_dia     = _num(get(row, 'SAL. DIA')),
            coste_hora  = _num(get(row, 'SAL. HORA')),
            dias_trab   = _num(get(row, 'TOTAL DIAS TRABAJADOS')),
            coste_total = _num(get(row, 'IMPORTE TOTAL')),
            atrasos     = _num(get(row, 'ATRASOS/OTROS')),
            obs         = _txt(get(row, 'OBSERVACIONES')),
            days        = day_hours,
        ))

    return workers, sum(w['COSTE TOTAL'] for w in workers)


def read_varias_obras(path):
    """Devuelve: (nombre_hoja, workers, total_coste_obra, total_trabajadores)"""
    todos, total_obra = _read_formato_diario(path)
    return ('Varias Obras', [w for w in todos if _incluir(w)], float(total_obra), len(todos))


def read_obra_xlsm(path, sheet_name):
    """
    Reporte de obra (.xlsm, hoja 'Reporte', cabecera en la fila 4).
    Si el fichero no tiene hoja 'Reporte' se lee con el formato de Varias Obras.
    Devuelve: (nombre_hoja, workers, total_coste_obra, total_trabajadores)
    """
    if 'Reporte' not in pd.ExcelFile(path).sheet_names:
        todos, total_obra = _read_formato_diario(path)
        for w in todos:
            w['OBRA'] = ''
        return (sheet_name, [w for w in todos if _incluir(w)], float(total_obra), len(todos))

    df = pd.read_excel(path, header=3, sheet_name='Reporte')
    df.columns = [_norm(c) for c in df.columns]
    df = df.dropna(subset=['OPERARIOS']).copy()
    df = df[~df['OPERARIOS'].astype(str).str.strip().str.upper().str.startswith('TOTAL')]

    # Columnas de horas extra por dia: "EXTRA 18/09/2026 (HORAS)"
    date_per_col = {}
    for col in df.columns:
        if col.startswith('EXTRA'):
            m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', col)
            if m:
                d, mo, y = m.groups()
                date_per_col[col] = f"{int(d):02d}-{int(mo):02d}-{y}"

    def get(row, *names):
        for n in names:
            if n in row.index:
                return row[n]
        return None

    workers = []
    for _, row in df.iterrows():
        day_hours = {}
        for col, fecha in date_per_col.items():
            v = _num(row[col])
            if v > 0:
                day_hours[fecha] = day_hours.get(fecha, 0) + v

        sal_dia    = _num(get(row, 'SALARIO X DIA'))
        coste_hora = _num(get(row, 'COSTE HORAS EXTRA')) or sal_dia / 8
        atrasos    = _num(get(row, 'ATRASOS/OTROS'))
        if 'COSTE TOTAL' in row.index:
            coste_total = _num(row['COSTE TOTAL'])
        else:
            coste_total = _num(get(row, 'TOTAL A PAGAR'))

        workers.append(_worker(
            nombre      = _txt(row['OPERARIOS']),
            ajustero    = _txt(get(row, 'AJUSTERO')),
            tecnico     = '',
            supervisor  = _txt(get(row, 'SUPERVISOR')),
            oficio      = _txt(get(row, 'OFICIO')),
            obra        = '',
            sal_dia     = sal_dia,
            coste_hora  = coste_hora,
            dias_trab   = _num(get(row, 'TOTAL DIAS ADMIN')) + _num(get(row, 'TOTAL DIAS AJUSTE')),
            coste_total = coste_total,
            atrasos     = atrasos,
            obs         = _txt(get(row, 'OBSERVACIONES')),
            days        = day_hours,
        ))

    total_obra = sum(w['COSTE TOTAL'] for w in workers)
    return (sheet_name, [w for w in workers if _incluir(w)], float(total_obra), len(workers))


def detect_obra_code(path):
    """Extrae el codigo de obra (5 digitos exactos) del nombre del fichero."""
    name = os.path.basename(str(path))
    nums = re.findall(r'(?<!\d)\d{5}(?!\d)', name)
    return nums[0] if nums else None


# ---------------------------------------------------------------------------
# Escritura del Excel
# ---------------------------------------------------------------------------

# Paleta azul — visualmente intuitiva, accesible para consulta rapida
HDR_BG   = "1F4E79"   # azul oscuro (cabecera)
HDR_FG   = "FFFFFF"
DAY_BG   = "BDD7EE"   # azul claro para cols de dias
TOT_BG   = "9DC3E6"   # azul medio para fila total
ALERT_FG = "C00000"   # rojo intenso para alertas (texto)

FMT_MONEY = '#,##0.00'
FMT_ACCT  = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
FMT_DEC1  = '0.0'
FMT_PCT   = '0.0%'

def _brd_outer():
    """Borde exterior completo."""
    s = Side(style='thin', color='999999')
    return Border(left=s, right=s, top=s, bottom=s)

def _brd_row():
    """Solo linea inferior (separador de filas)."""
    b = Side(style='thin', color='CCCCCC')
    n = Side(style=None)
    return Border(bottom=b, left=n, right=n, top=n)

def _hdr_font():  return Font(name='Calibri', bold=True, color=HDR_FG, size=10)
def _hdr_fill():  return PatternFill('solid', start_color=HDR_BG)
def _day_fill():  return PatternFill('solid', start_color=DAY_BG)
def _tot_fill():  return PatternFill('solid', start_color=TOT_BG)
def _cen():       return Alignment(horizontal='center', vertical='center')
def _left():      return Alignment(horizontal='left',   vertical='center')


TEXT_COLS = ('TRABAJADOR', 'AJUSTERO', 'TECNICO', 'SUPERVISOR', 'OFICIO', 'OBRA', 'COMENTARIOS')

LABELS = {
    'TRABAJADOR':     'Trabajador',
    'AJUSTERO':       'Ajustero',
    'TECNICO':        'Técnico Obra',
    'SUPERVISOR':     'Supervisor',
    'OFICIO':         'Oficio',
    'OBRA':           'Obra',
    'SAL. DIA':       'Sal. Dia',
    'COSTE HORA':     'Coste/Hora',
    'DIAS':           'Dias quincena trabajados',
    'COSTE DIAS':     'Coste dias / ajuste',
    'COSTE H. EXTRA': 'Coste H. Extra',
    'TOTAL H. EXTRA': 'Total H. Extra',
    '% H. EXTRA':     '% H. Extra',
    'COSTE TOTAL':    'Coste total quincena',
    'DIAS ATR':       'Dias atrasados',
    'HORAS ATR':      'Horas extras atrasadas',
    'ATRASOS':        'Atrasos/otros',
    'NOCTURNO':       'Coste nocturno',
    'FESTIVO':        'Dia festivo',
    'COSTE FINAL':    'coste total',
    'COMENTARIOS':    'comentarios',
}

COL_WIDTHS = {
    'TRABAJADOR': 30, 'AJUSTERO': 26, 'TECNICO': 18, 'SUPERVISOR': 20, 'OFICIO': 14,
    'OBRA': 14, 'SAL. DIA': 10, 'COSTE HORA': 11, 'DIAS': 25, 'COSTE DIAS': 22,
    'COSTE H. EXTRA': 14, 'TOTAL H. EXTRA': 13, '% H. EXTRA': 12, 'COSTE TOTAL': 20,
    'DIAS ATR': 14, 'HORAS ATR': 21, 'ATRASOS': 16, 'NOCTURNO': 15, 'FESTIVO': 13,
    'COSTE FINAL': 16, 'COMENTARIOS': 45,
}

# Columnas que llevan SUM en la fila de totales
SUM_COLS = ('DIAS', 'COSTE DIAS', 'COSTE H. EXTRA', 'TOTAL H. EXTRA', 'COSTE TOTAL',
            'ATRASOS', 'NOCTURNO', 'FESTIVO', 'COSTE FINAL')


def write_sheet(wb, sheet_name, workers, varias=False):
    """
    Escribe la hoja de un reporte. Devuelve la fila de totales (o None si no hay
    trabajadores) y el mapa columna -> letra, para enlazar desde el Resumen.
    """
    ws = wb.create_sheet(title=sheet_name[:31])

    if not workers:
        ws['A1'] = 'Sin trabajadores con horas extra ni atrasos'
        return None, {}

    all_days = sorted({d for w in workers for d in w['_days']}, key=_fecha_key)

    fixed = ['TRABAJADOR', 'AJUSTERO']
    for key in ('TECNICO', 'SUPERVISOR', 'OFICIO', 'OBRA'):
        if any(w.get(key) for w in workers):
            fixed.append(key)
    fixed += ['SAL. DIA', 'COSTE HORA']

    tail = ['DIAS', 'COSTE DIAS', 'COSTE H. EXTRA', 'TOTAL H. EXTRA', '% H. EXTRA',
            'COSTE TOTAL', 'DIAS ATR', 'HORAS ATR', 'ATRASOS']
    if varias:
        tail += ['NOCTURNO', 'FESTIVO']
    tail += ['COSTE FINAL', 'COMENTARIOS']

    all_cols = fixed + all_days + tail
    L = {col: get_column_letter(i) for i, col in enumerate(all_cols, 1)}

    # Umbral dinamico de horas extra (sobre trabajadores con horas extra):
    #   - al menos 2.3x la media Y al menos 25h por encima de la media
    con_he     = [w for w in workers if w['TOTAL H. EXTRA'] > 0]
    horas      = [w['TOTAL H. EXTRA'] for w in con_he]
    media      = sum(horas) / len(horas) if horas else 0
    umbral_h   = max(media * 2.3, media + 25)

    # Umbral dinamico de % H. Extra sobre coste total:
    #   - al menos 1.5x el % medio Y 10 puntos por encima, nunca bajo el minimo
    pcts       = [w['COSTE H. EXTRA'] / w['COSTE TOTAL'] for w in con_he if w['COSTE TOTAL'] > 0]
    media_pct  = sum(pcts) / len(pcts) if pcts else 0
    umbral_pct = max(media_pct * 1.5, media_pct + 0.10, CONFIG_PORCENTAJE_MINIMO)

    # ── Fila 1: cabecera ─────────────────────────────────────────────────────
    for ci, col in enumerate(all_cols, 1):
        c = ws.cell(1, ci, LABELS.get(col, col))
        if col in all_days:
            c.fill = _day_fill()
            c.font = Font(name='Calibri', bold=True, color=HDR_BG, size=9)
        else:
            c.fill = _hdr_fill()
            c.font = _hdr_font()
        c.alignment = _cen()
        c.border    = _brd_outer()
    ws.row_dimensions[1].height = 20

    # ── Filas de datos ───────────────────────────────────────────────────────
    first_day, last_day = (L[all_days[0]], L[all_days[-1]]) if all_days else (None, None)

    for r, w in enumerate(workers, 2):
        coste_total = w['COSTE TOTAL']
        pct_extra   = (w['COSTE H. EXTRA'] / coste_total) if coste_total > 0 else 0
        es_alerta   = w['TOTAL H. EXTRA'] > umbral_h or (pct_extra > umbral_pct and pct_extra > 0)

        formulas = {
            'COSTE DIAS':     f"={L['COSTE TOTAL']}{r}-{L['COSTE H. EXTRA']}{r}",
            'COSTE H. EXTRA': f"={L['TOTAL H. EXTRA']}{r}*{L['COSTE HORA']}{r}",
            'TOTAL H. EXTRA': f"=SUM({first_day}{r}:{last_day}{r})" if all_days else 0,
            '% H. EXTRA':     f"=IFERROR({L['COSTE H. EXTRA']}{r}/{L['COSTE TOTAL']}{r},0)",
            'ATRASOS':        f"={L['DIAS ATR']}{r}*{L['SAL. DIA']}{r}+{L['COSTE HORA']}{r}*{L['HORAS ATR']}{r}",
            'COSTE FINAL':    '=' + '+'.join(f"N({L[k]}{r})" for k in
                                             ('COSTE H. EXTRA', 'ATRASOS', 'NOCTURNO', 'FESTIVO') if k in L),
        }

        for ci, col in enumerate(all_cols, 1):
            if col in formulas:
                val = formulas[col]
            elif col in all_days:
                val = w['_days'].get(col)
            elif col in ('NOCTURNO', 'FESTIVO'):
                val = None
            else:
                val = w.get(col)
            if val == '' or (isinstance(val, float) and val == 0.0 and col in all_days):
                val = None

            bold  = (col in ('TOTAL H. EXTRA', '% H. EXTRA') and es_alerta) or col == 'COSTE FINAL'
            color = ALERT_FG if es_alerta else '000000'
            if col in all_days and val is not None and val > CONFIG_HORAS_DIA_ALERTA:
                color, bold = ALERT_FG, True

            c = ws.cell(r, ci, val)
            c.font      = Font(name='Calibri', size=10, bold=bold, color=color)
            c.border    = _brd_row()
            c.alignment = _left() if col in TEXT_COLS else _cen()
            if col in ('SAL. DIA', 'COSTE HORA', 'COSTE H. EXTRA', 'COSTE TOTAL'):
                c.number_format = FMT_MONEY
            elif col in ('COSTE DIAS', 'ATRASOS', 'NOCTURNO', 'FESTIVO', 'COSTE FINAL'):
                c.number_format = FMT_ACCT
            elif col == 'TOTAL H. EXTRA' or col in all_days:
                c.number_format = FMT_DEC1
            elif col == '% H. EXTRA':
                c.number_format = FMT_PCT

    # ── Fila totales ─────────────────────────────────────────────────────────
    last_data = 1 + len(workers)
    tr        = last_data + 1
    for ci, col in enumerate(all_cols, 1):
        c = ws.cell(tr, ci)
        c.fill      = _tot_fill()
        c.font      = Font(name='Calibri', bold=True, size=10)
        c.border    = _brd_outer()
        c.alignment = _left() if ci == 1 else _cen()
        cl = L[col]
        if ci == 1:
            c.value = 'Total'
        elif col in all_days or col in SUM_COLS:
            c.value = f'=SUM({cl}2:{cl}{last_data})'
            if col in all_days or col == 'TOTAL H. EXTRA':
                c.number_format = FMT_DEC1
            elif col in ('COSTE H. EXTRA', 'COSTE TOTAL'):
                c.number_format = FMT_MONEY
            elif col != 'DIAS':
                c.number_format = FMT_ACCT
        elif col == '% H. EXTRA':
            c.value         = f"=IFERROR({L['COSTE H. EXTRA']}{tr}/{L['COSTE TOTAL']}{tr},0)"
            c.number_format = FMT_PCT

    # ── Anchos, filtros y paneles ────────────────────────────────────────────
    for ci, col in enumerate(all_cols, 1):
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTHS.get(col, 10)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(all_cols))}{last_data}"
    ws.freeze_panes = 'A2'
    return tr, L


def _ref(sheet_title, cell):
    """Referencia a celda de otra hoja: 'Varias Obras'!Y45"""
    return "'" + sheet_title.replace("'", "''") + "'!" + cell


def build_excel(sheets, output_path):
    """sheets es lista de (sheet_name, workers, coste_total_obra, total_trabajadores)."""
    wb = Workbook()
    wb.remove(wb.active)

    # ── Hoja resumen ──────────────────────────────────────────────────────────
    # Obra | Trabajadores | Trab. con HE | Coste Total | Coste H. Extra + atrasos
    #      | Total H. Extra | %   — enlazado con las hojas de cada reporte, de
    # forma que al rellenar atrasos a mano el resumen se actualiza solo.
    ws_r = wb.create_sheet('Resumen', 0)

    headers = ['Obra', 'Trabajadores', 'Trab. con HE', 'Coste Total',
               'Coste H. Extra + atrasos', 'Total H. Extra', '% s/Total']
    for ci, label in enumerate(headers, 1):
        c = ws_r.cell(1, ci, label)
        c.font      = _hdr_font()
        c.fill      = _hdr_fill()
        c.alignment = _cen()
        c.border    = _brd_outer()
    ws_r.row_dimensions[1].height = 18

    for ri, (sname, workers, coste_total_obra, total_trab) in enumerate(sheets, 2):
        tr, L = write_sheet(wb, sname, workers, varias=(sname == 'Varias Obras'))
        title = sname[:31]
        if tr:
            last = tr - 1
            con_he   = f"=COUNTIF({_ref(title, L['TOTAL H. EXTRA'] + '2')}:{L['TOTAL H. EXTRA']}{last},\">0\")"
            coste_he = '=' + _ref(title, f"{L['COSTE FINAL']}{tr}")
            total_he = '=' + _ref(title, f"{L['TOTAL H. EXTRA']}{tr}")
        else:
            con_he = coste_he = total_he = 0

        valores = [sname, total_trab, con_he, coste_total_obra, coste_he, total_he,
                   f'=IFERROR(E{ri}/D{ri},0)']
        for ci, val in enumerate(valores, 1):
            c = ws_r.cell(ri, ci, val)
            c.font      = Font(name='Calibri', size=10)
            c.border    = _brd_row()
            c.alignment = _left() if ci == 1 else _cen()
            if ci in (4, 5): c.number_format = FMT_MONEY
            if ci == 6:      c.number_format = FMT_DEC1
            if ci == 7:      c.number_format = FMT_PCT

    tr = 1 + len(sheets) + 1
    for ci in range(1, 8):
        c = ws_r.cell(tr, ci)
        c.fill      = _tot_fill()
        c.font      = Font(name='Calibri', bold=True, size=10)
        c.border    = _brd_outer()
        c.alignment = _left() if ci == 1 else _cen()
        cl = get_column_letter(ci)
        if ci == 1:
            c.value = 'Total'
        elif ci < 7:
            c.value = f'=SUM({cl}2:{cl}{tr-1})'
            if ci in (4, 5): c.number_format = FMT_MONEY
            if ci == 6:      c.number_format = FMT_DEC1
        else:
            # % global = total coste extras + atrasos / total coste obras
            c.value = f'=IFERROR(E{tr}/D{tr},0)'
            c.number_format = FMT_PCT

    for col, width in zip('ABCDEFG', (22, 13, 13, 16, 24, 15, 12)):
        ws_r.column_dimensions[col].width = width
    ws_r.auto_filter.ref = f"A1:G{tr-1}"
    ws_r.freeze_panes = 'A2'

    wb.save(output_path)


# ---------------------------------------------------------------------------
# Interfaz grafica
# ---------------------------------------------------------------------------




def resource_path(filename):
    """Devuelve la ruta al recurso, compatible con PyInstaller (--onefile)."""
    import os
    base = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, filename)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Reporte de Horas Extra")
        self.resizable(False, False)
        self.configure(bg="#f0f4f8")
        try:
            self.iconbitmap(resource_path("logo_programa_multi.ico"))
        except Exception:
            pass
        # Estado
        self._varias_path  = None  # str o None
        self._varias_row   = None  # widget de la fila
        self._obras        = []    # lista de (path, sheet_label, widget_frame)
        # UI
        self._build_ui()
        self.eval("tk::PlaceWindow . center")

    # ── Construcción de la UI ────────────────────────────────────────────────
    def _build_ui(self):
        tk.Label(self, text="Reporte de Horas Extra",
                 font=("Arial", 14, "bold"), bg="#1F4E79", fg="white",
                 pady=10, padx=20).grid(row=0, column=0, sticky="ew")

        main = tk.Frame(self, bg="#f0f4f8", padx=16, pady=12)
        main.grid(row=1, column=0, sticky="nsew")
        main.columnconfigure(0, weight=1)

        # ── Sección Varias Obras ─────────────────────────────────────────────
        sec_varias = tk.LabelFrame(main, text=" Reporte de Varias Obras ",
                                   font=("Arial", 10, "bold"),
                                   fg="#1F4E79", bg="#f0f4f8",
                                   padx=10, pady=8)
        sec_varias.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        sec_varias.columnconfigure(0, weight=1)

        self._frame_varias = tk.Frame(sec_varias, bg="#f0f4f8")
        self._frame_varias.grid(row=0, column=0, sticky="ew")
        self._frame_varias.columnconfigure(0, weight=1)

        self._lbl_empty_varias = tk.Label(self._frame_varias,
            text="Aún no se ha añadido el reporte de Varias Obras.",
            font=("Arial", 9, "italic"), bg="#f0f4f8", fg="#888")
        self._lbl_empty_varias.grid(row=0, column=0, pady=4, sticky="w")

        self._btn_add_varias = tk.Button(sec_varias,
                  text="+  Añadir reporte de Varias Obras",
                  font=("Arial", 9), relief="flat",
                  bg="#dce6f1", cursor="hand2",
                  command=self._add_varias)
        self._btn_add_varias.grid(row=1, column=0, sticky="w", pady=(8, 0))

        # ── Sección Obras adicionales ────────────────────────────────────────
        sec_obras = tk.LabelFrame(main, text=" Reportes de obra ",
                                  font=("Arial", 10, "bold"),
                                  fg="#1F4E79", bg="#f0f4f8",
                                  padx=10, pady=8)
        sec_obras.grid(row=1, column=0, sticky="ew", pady=(0, 10))
        sec_obras.columnconfigure(0, weight=1)

        self._frame_obras = tk.Frame(sec_obras, bg="#f0f4f8")
        self._frame_obras.grid(row=0, column=0, sticky="ew")
        self._frame_obras.columnconfigure(0, weight=1)

        self._lbl_empty_obras = tk.Label(self._frame_obras,
            text="Aún no se ha añadido ningún reporte de obra.",
            font=("Arial", 9, "italic"), bg="#f0f4f8", fg="#888")
        self._lbl_empty_obras.grid(row=0, column=0, pady=4, sticky="w")

        tk.Button(sec_obras, text="+  Añadir reporte(s) de obra",
                  font=("Arial", 9), relief="flat",
                  bg="#dce6f1", cursor="hand2",
                  command=self._add_obras
                  ).grid(row=1, column=0, sticky="w", pady=(8, 0))

        # ── Botones Generar + Ajustes ────────────────────────────────────────
        btn_row = tk.Frame(main, bg="#f0f4f8")
        btn_row.grid(row=2, column=0, pady=(4, 4))

        tk.Button(btn_row, text="▶  Generar informe", font=("Arial", 11, "bold"),
                  bg="#4CAF50", fg="white", activebackground="#388E3C",
                  relief="flat", cursor="hand2", padx=18, pady=8,
                  command=self._run
                  ).grid(row=0, column=0, padx=(0, 8))

        tk.Button(btn_row, text="⚙  Ajustes", font=("Arial", 10),
                  bg="#dce6f1", relief="flat", cursor="hand2",
                  padx=12, pady=8,
                  command=self._open_settings
                  ).grid(row=0, column=1)

        self.progress = ttk.Progressbar(main, mode="indeterminate", length=420)
        self.progress.grid(row=3, column=0, sticky="ew", pady=(0, 2))

        self.lbl_status = tk.Label(main, text="", font=("Arial", 9),
                                   bg="#f0f4f8", fg="#444")
        self.lbl_status.grid(row=4, column=0, pady=(0, 4))

    # ── Helper: crear fila visual de reporte cargado ─────────────────────────
    def _make_row(self, parent, label_text, path, on_remove):
        row = tk.Frame(parent, bg="#f0f4f8")
        row.columnconfigure(2, weight=1)
        tk.Label(row, text="OK", font=("Arial", 9, "bold"),
                 bg="#4CAF50", fg="white", width=3, pady=2
                 ).grid(row=0, column=0, padx=(0, 8))
        tk.Label(row, text=label_text, font=("Arial", 10, "bold"),
                 bg="#f0f4f8", fg="#1F4E79", width=12, anchor="w"
                 ).grid(row=0, column=1, padx=(0, 6), sticky="w")
        short = str(path).split("/")[-1].split("\\")[-1]
        tk.Label(row, text=short, font=("Arial", 9),
                 bg="#f0f4f8", fg="#666", anchor="w"
                 ).grid(row=0, column=2, sticky="w")
        tk.Button(row, text="✕", font=("Arial", 9, "bold"),
                  relief="flat", bg="#f0f4f8", fg="#c00", cursor="hand2",
                  command=on_remove
                  ).grid(row=0, column=3, padx=(6, 0))
        return row

    # ── Acciones: Varias Obras ───────────────────────────────────────────────
    def _add_varias(self):
        p = filedialog.askopenfilename(
            title="Selecciona el reporte de Varias Obras",
            filetypes=[("Excel", "*.xlsx *.xlsm *.xls")])
        if not p:
            return
        self._varias_path = p
        if self._lbl_empty_varias.winfo_exists():
            self._lbl_empty_varias.grid_remove()
        self._varias_row = self._make_row(self._frame_varias,
                                          "Varias Obras", p,
                                          on_remove=self._remove_varias)
        self._varias_row.grid(row=1, column=0, sticky="ew", pady=2)
        self._btn_add_varias.config(state="disabled")
        self.update_idletasks()
        self.geometry("")

    def _remove_varias(self):
        if self._varias_row is not None:
            self._varias_row.destroy()
            self._varias_row = None
        self._varias_path = None
        self._lbl_empty_varias.grid()
        self._btn_add_varias.config(state="normal")
        self.update_idletasks()
        self.geometry("")

    # ── Acciones: Reportes de obra (multiselección) ──────────────────────────
    def _add_obras(self):
        paths = filedialog.askopenfilenames(
            title="Selecciona uno o varios reportes de obra",
            filetypes=[("Excel", "*.xlsx *.xlsm *.xls")])
        if not paths:
            return
        if self._lbl_empty_obras.winfo_exists():
            self._lbl_empty_obras.grid_remove()
        for p in paths:
            # Evitar duplicados (si el usuario añade el mismo archivo dos veces)
            if any(existing == p for existing, _, _ in self._obras):
                continue
            cod = detect_obra_code(p)
            sheet_label = cod if cod else f"Reporte {len(self._obras) + 1}"
            row = self._make_row(self._frame_obras, sheet_label, p,
                                 on_remove=lambda pp=p: self._remove_obra(pp))
            row.grid(row=len(self._obras), column=0, sticky="ew", pady=2)
            self._obras.append((p, sheet_label, row))
        self.update_idletasks()
        self.geometry("")

    def _remove_obra(self, path):
        for i, (p, _, frame) in enumerate(self._obras):
            if p == path:
                frame.destroy()
                self._obras.pop(i)
                break
        for i, (_, _, frame) in enumerate(self._obras):
            frame.grid_configure(row=i)
        if not self._obras:
            self._lbl_empty_obras.grid()
        self.update_idletasks()
        self.geometry("")

    # ── Diálogo de configuración ─────────────────────────────────────────────
    def _open_settings(self):
        global CONFIG_HORAS_DIA_ALERTA, CONFIG_PORCENTAJE_MINIMO

        dlg = tk.Toplevel(self)
        dlg.title("Ajustes de umbrales")
        dlg.resizable(False, False)
        dlg.configure(bg="#f0f4f8")
        dlg.transient(self)
        dlg.grab_set()
        try:
            dlg.iconbitmap(resource_path("logo_programa_multi.ico"))
        except Exception:
            pass

        tk.Label(dlg, text="Configurar umbrales de alerta",
                 font=("Arial", 12, "bold"), bg="#1F4E79", fg="white",
                 pady=8, padx=16).grid(row=0, column=0, columnspan=2, sticky="ew")

        body = tk.Frame(dlg, bg="#f0f4f8", padx=18, pady=14)
        body.grid(row=1, column=0, columnspan=2, sticky="ew")

        # ── Horas por día ───────────────────────────────────────────────────
        tk.Label(body,
                 text="Horas en un día para resaltar la celda (en rojo/negrita):",
                 font=("Arial", 10), bg="#f0f4f8", anchor="w"
                 ).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 4))

        var_horas = tk.IntVar(value=CONFIG_HORAS_DIA_ALERTA)
        tk.Spinbox(body, from_=1, to=12, textvariable=var_horas,
                   width=6, font=("Arial", 10),
                   relief="solid", bd=1
                   ).grid(row=1, column=0, sticky="w", padx=(0, 6))
        tk.Label(body, text="horas", font=("Arial", 9, "italic"),
                 bg="#f0f4f8", fg="#666"
                 ).grid(row=1, column=1, sticky="w")

        # ── Porcentaje ─────────────────────────────────────────────────────
        tk.Label(body,
                 text="Porcentaje mínimo de H. Extra sobre coste total para marcar fila:",
                 font=("Arial", 10), bg="#f0f4f8", anchor="w"
                 ).grid(row=2, column=0, columnspan=2, sticky="w", pady=(14, 4))

        var_pct = tk.IntVar(value=int(round(CONFIG_PORCENTAJE_MINIMO * 100)))
        tk.Spinbox(body, from_=5, to=100, increment=5, textvariable=var_pct,
                   width=6, font=("Arial", 10),
                   relief="solid", bd=1
                   ).grid(row=3, column=0, sticky="w", padx=(0, 6))
        tk.Label(body, text="%", font=("Arial", 9, "italic"),
                 bg="#f0f4f8", fg="#666"
                 ).grid(row=3, column=1, sticky="w")

        # ── Botones ────────────────────────────────────────────────────────
        btns = tk.Frame(dlg, bg="#f0f4f8", pady=10, padx=18)
        btns.grid(row=2, column=0, columnspan=2, sticky="ew")

        def _save():
            global CONFIG_HORAS_DIA_ALERTA, CONFIG_PORCENTAJE_MINIMO
            try:
                h = int(var_horas.get())
                p = float(var_pct.get()) / 100.0
                if not (1 <= h <= 12):
                    raise ValueError("Horas fuera de rango (1-12)")
                if not (0.05 <= p <= 1.0):
                    raise ValueError("Porcentaje fuera de rango (5-100)")
            except Exception as e:
                messagebox.showerror("Valor no válido",
                                     f"Comprueba los valores introducidos.\n\n{e}",
                                     parent=dlg)
                return
            CONFIG_HORAS_DIA_ALERTA  = h
            CONFIG_PORCENTAJE_MINIMO = p
            guardar_configuracion()
            dlg.destroy()
            self.lbl_status.config(
                text=f"Ajustes guardados: {h}h/día · {int(p*100)}% del coste",
                fg="#1a6b1a")

        def _cancel():
            dlg.destroy()

        tk.Button(btns, text="Guardar", font=("Arial", 10, "bold"),
                  bg="#4CAF50", fg="white", activebackground="#388E3C",
                  relief="flat", cursor="hand2", padx=16, pady=6,
                  command=_save
                  ).pack(side="right", padx=(6, 0))
        tk.Button(btns, text="Cancelar", font=("Arial", 10),
                  bg="#e0e0e0", relief="flat", cursor="hand2",
                  padx=14, pady=6,
                  command=_cancel
                  ).pack(side="right")

        # Centrar la ventana sobre la principal
        dlg.update_idletasks()
        x = self.winfo_rootx() + (self.winfo_width()  - dlg.winfo_width())  // 2
        y = self.winfo_rooty() + (self.winfo_height() - dlg.winfo_height()) // 2
        dlg.geometry(f"+{x}+{y}")

    # ── Generación del informe ───────────────────────────────────────────────
    def _run(self):
        if not self._varias_path and not self._obras:
            messagebox.showwarning(
                "Faltan reportes",
                "Selecciona al menos un reporte antes de generar el informe.")
            return

        output = filedialog.asksaveasfilename(
            title="Guardar informe de horas extra y atrasos como…",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile="Reporte_Horas_Extras_Atrasos.xlsx")
        if not output:
            return

        self.progress.start(10)
        self.lbl_status.config(text="Procesando…", fg="#444")
        self.update()

        try:
            sheets = []
            if self._varias_path:
                name, workers, total_obra, total_trab = read_varias_obras(self._varias_path)
                sheets.append((name, workers, total_obra, total_trab))
            for path, sheet_label, _ in self._obras:
                name, workers, total_obra, total_trab = read_obra_xlsm(path, sheet_label)
                sheets.append((name, workers, total_obra, total_trab))

            build_excel(sheets, output)

            total_w  = sum(len(w) for _, w, _, _ in sheets)
            total_he = sum(x['TOTAL H. EXTRA'] for _, w, _, _ in sheets for x in w)

            self.lbl_status.config(
                text=f"✔ Listo — {len(sheets)} hojas · {total_w} trabajadores · {total_he:.1f} h extra",
                fg="#1a6b1a")

            if messagebox.askyesno(
                "Informe generado",
                f"El informe se ha guardado en:\n{output}\n\n¿Quieres abrirlo ahora?"):
                import os, subprocess, platform
                if platform.system() == "Windows":
                    os.startfile(output)
                elif platform.system() == "Darwin":
                    subprocess.call(["open", output])
                else:
                    subprocess.call(["xdg-open", output])

        except Exception as e:
            import traceback
            messagebox.showerror(
                "Error",
                f"Ha ocurrido un error al procesar los reportes:\n\n{e}\n\n{traceback.format_exc()}")
            self.lbl_status.config(text="Error al procesar.", fg="red")
        finally:
            self.progress.stop()


# ---------------------------------------------------------------------------
if __name__ == "__main__":
    cargar_configuracion()
    app = App()
    app.mainloop()
